#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""工业门报价单生成器（CSV/JSON → 格式化 Excel 报价单）

用法:
    python generate_quote.py                    # 使用本目录 items.csv + config.json
    python generate_quote.py --items a.csv --config c.json --out 报价单.xlsx
    python generate_quote.py --demo             # 生成演示输入文件（config.example.json / items.example.csv）

依赖: openpyxl   →  pip install openpyxl
生成后用 Excel / WPS / LibreOffice 打开并「导出 PDF」即可发给客户。
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖：请先安装 openpyxl（pip install openpyxl）")

DEFAULT_CONFIG = {
    "company": "示例门业有限公司",
    "title": "产品报价单",
    "contact": "联系人：张三    电话：138-0000-0000",
    "currency": "¥",
    "tax_rate": 0.13,
    "tax_label": "增值税(13%)",
    "note": (
        "1. 本报价为含税价，运杂费、安装费另计（如含安装将在备注注明）；\n"
        "2. 报价有效期 15 天；\n"
        "3. 付款方式：合同签订预付 50%，发货前付清尾款。"
    ),
    "columns": ["序号", "项目名称", "规格型号", "单位", "数量", "单价", "小计", "备注"],
    "widths": [6, 22, 26, 8, 8, 12, 14, 18],
    "font_name": "微软雅黑",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_config(path: Path) -> dict:
    cfg = dict(DEFAULT_CONFIG)
    if path.exists():
        user = json.loads(path.read_text(encoding="utf-8"))
        cfg.update(user)
    return cfg


def load_items(path: Path) -> list:
    """支持 .csv（表头：项目名称,规格型号,单位,数量,单价,备注）与 .json（对象数组）。"""
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        return [dict(x) for x in data]
    for enc in ("utf-8-sig", "gb18030"):
        try:
            with path.open(newline="", encoding=enc) as f:
                return [dict(row) for row in csv.DictReader(f)]
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法解码 CSV 文件（尝试过 utf-8-sig / gb18030）：{path}")


def num(x, default=0.0):
    try:
        return float(str(x).replace(",", "").strip())
    except (TypeError, ValueError):
        return default


def pick(item: dict, cn_key: str, en_key: str, default=""):
    v = item.get(cn_key)
    if v in (None, ""):
        v = item.get(en_key)
    return default if v is None else v


def build_quote(cfg: dict, items: list, out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"
    font = cfg["font_name"]
    ncols = len(cfg["columns"])

    # 标题行
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(1, 1, cfg["title"])
    t.font = Font(name=font, size=18, bold=True)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # 公司 / 日期 / 联系方式
    split = max(4, ncols // 2)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=split)
    c = ws.cell(2, 1, cfg["company"])
    c.font = Font(name=font, size=12, bold=True)
    ws.merge_cells(start_row=2, start_column=split + 1, end_row=2, end_column=ncols)
    d = ws.cell(2, split + 1, f"日期：{date.today():%Y-%m-%d}")
    d.font = Font(name=font, size=10)
    d.alignment = Alignment(horizontal="right")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    ct = ws.cell(3, 1, cfg["contact"])
    ct.font = Font(name=font, size=10, color="555555")

    # 表头
    hr = 4
    for j, name in enumerate(cfg["columns"], 1):
        cell = ws.cell(hr, j, name)
        cell.font = Font(name=font, size=11, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    # 明细
    subtotal = 0.0
    r = hr
    for i, item in enumerate(items, 1):
        r = hr + i
        qty = num(pick(item, "数量", "qty"))
        price = num(pick(item, "单价", "price"))
        amount = qty * price
        subtotal += amount
        vals = [
            i,
            pick(item, "项目名称", "name"),
            pick(item, "规格型号", "spec"),
            pick(item, "单位", "unit"),
            qty,
            price,
            amount,
            pick(item, "备注", "remark"),
        ]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(r, j, v)
            cell.font = Font(name=font, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center" if j in (1, 4, 5) else ("right" if j in (6, 7) else "left"),
                vertical="center",
            )
            if j in (6, 7):
                cell.number_format = "#,##0.00"
    last = r if items else hr

    # 合计区
    cur = cfg["currency"]
    tax = subtotal * float(cfg.get("tax_rate") or 0)
    total = subtotal + tax
    rr = last + 1
    for label, value in (("合计", subtotal), (cfg.get("tax_label", "税金"), tax), ("价税总计", total)):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=ncols - 2)
        lc = ws.cell(rr, 1, label)
        lc.font = Font(name=font, size=11, bold=True)
        lc.alignment = Alignment(horizontal="right", vertical="center")
        vc = ws.cell(rr, ncols - 1, value)
        vc.font = Font(name=font, size=11, bold=True)
        vc.number_format = f'"{cur}"#,##0.00'
        vc.alignment = Alignment(horizontal="right", vertical="center")
        for j in range(1, ncols + 1):
            ws.cell(rr, j).border = BORDER
        rr += 1

    # 备注
    rr += 1
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=ncols)
    nc = ws.cell(rr, 1, cfg["note"])
    nc.font = Font(name=font, size=9, color="666666")
    nc.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
    ws.row_dimensions[rr].height = 52

    # 列宽
    for j, w in enumerate(cfg["widths"], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    wb.save(out_path)


def write_demo(dirpath: Path) -> None:
    (dirpath / "config.example.json").write_text(
        json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (dirpath / "items.example.csv").write_text(
        "项目名称,规格型号,单位,数量,单价,备注\n"
        "快速卷帘门,3.0m×3.0m PVC 快速门,樘,2,6800,含安装调试\n"
        "工业滑升门,4.0m×4.5m 保温滑升门,樘,1,15800,含电机\n"
        "硬质快速门,3.5m×3.5m 铝合金硬质,樘,1,26000,\n",
        encoding="utf-8-sig",
    )
    print(f"已生成演示文件到 {dirpath}（config.example.json / items.example.csv）")


def main() -> None:
    ap = argparse.ArgumentParser(description="工业门报价单生成器（CSV/JSON → Excel）")
    ap.add_argument("--items", default="items.csv", help="明细文件（.csv 或 .json）")
    ap.add_argument("--config", default="config.json", help="配置文件（不存在则用内置默认值）")
    ap.add_argument("--out", default=None, help="输出文件名（默认 报价单_YYYYMMDD.xlsx）")
    ap.add_argument("--demo", action="store_true", help="在脚本目录生成演示文件后退出")
    args = ap.parse_args()

    base = Path(__file__).resolve().parent
    if args.demo:
        write_demo(base)
        return

    cfg_path = Path(args.config)
    if not cfg_path.is_absolute():
        cfg_path = base / cfg_path
    cfg = load_config(cfg_path)

    items_path = Path(args.items)
    if not items_path.is_absolute():
        items_path = base / items_path
    if not items_path.exists():
        sys.exit(f"找不到明细文件：{items_path}（可用 --demo 生成演示文件）")

    items = load_items(items_path)
    out = Path(args.out) if args.out else base / f"报价单_{date.today():%Y%m%d}.xlsx"
    build_quote(cfg, items, out)
    print(f"已生成：{out}（{len(items)} 条明细）")


if __name__ == "__main__":
    main()
